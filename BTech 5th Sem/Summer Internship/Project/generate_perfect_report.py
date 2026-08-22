import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc):
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = "Page "
        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        p._p.append(fldSimple)

def main():
    doc = Document()
    
    # Page Margins
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.85)
    
    # Fonts and styling definitions
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    
    def add_standard_p(text, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
        p = doc.add_paragraph()
        p.alignment = alignment
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        
        # Check if it has a bold title prefix (like "1.1 Background")
        if (text.startswith("1.1") or text.startswith("1.2") or text.startswith("1.3") or text.startswith("1.4") or 
            text.startswith("1.5") or text.startswith("2.1") or text.startswith("2.2") or text.startswith("3.1") or 
            text.startswith("3.2") or text.startswith("4.1") or text.startswith("4.2") or text.startswith("5.1") or 
            text.startswith("5.2") or text.startswith("5.3") or text.startswith("5.4") or text.startswith("5.5") or 
            text.startswith("6.1") or text.startswith("6.2") or text.startswith("6.3") or text.startswith("6.4") or 
            text.startswith("6.5") or text.startswith("6.6") or text.startswith("7.1") or text.startswith("7.2") or 
            text.startswith("7.3") or text.startswith("7.4") or text.startswith("7.5") or text.startswith("8.1") or 
            text.startswith("8.2") or text.startswith("8.3") or text.startswith("8.4") or text.startswith("9.1") or 
            text.startswith("9.2") or text.startswith("9.3") or text.startswith("9.4") or text.startswith("9.5") or 
            text.startswith("9.6") or text.startswith("9.7") or text.startswith("9.8") or text.startswith("9.9") or 
            text.startswith("9.10") or text.startswith("10.1") or text.startswith("10.2") or text.startswith("10.3") or 
            text.startswith("10.4") or text.startswith("10.5") or text.startswith("10.6") or text.startswith("11.1") or 
            text.startswith("11.2") or text.startswith("11.3") or text.startswith("11.4") or text.startswith("12.1") or 
            text.startswith("12.2") or text.startswith("12.3") or text.startswith("13.1") or text.startswith("14.1") or 
            text.startswith("14.2") or text.startswith("14.3") or text.startswith("14.4") or text.startswith("14.5") or 
            text.startswith("14.6") or text.startswith("14.7") or text.startswith("14.8") or text.startswith("14.9") or 
            text.startswith("14.10") or text.startswith("14.11") or text.startswith("14.12") or text.startswith("14.13") or 
            text.startswith("14.14") or text.startswith("14.15") or text.startswith("14.16") or text.startswith("15.1") or 
            text.startswith("15.2") or text.startswith("15.3") or text.startswith("15.4") or text.startswith("16.1") or 
            text.startswith("16.2") or text.startswith("16.3") or text.startswith("16.4") or text.startswith("16.5") or 
            text.startswith("16.6") or text.startswith("16.7") or text.startswith("16.8") or text.startswith("16.9") or 
            text.startswith("16.10") or text.startswith("17.1") or text.startswith("17.2") or text.startswith("17.3") or 
            text.startswith("17.4") or text.startswith("17.5") or text.startswith("18.1") or text.startswith("19.1") or 
            text.startswith("20.1") or text.startswith("21.1")):
            
            parts = text.split("\n", 1)
            p_hdr = p.add_run(parts[0] + "\n")
            p_hdr.font.name = 'Times New Roman'
            p_hdr.font.bold = True
            p_hdr.font.size = Pt(13)
            p_hdr.font.color.rgb = RGBColor(44, 82, 130)
            
            if len(parts) > 1:
                p_body = p.add_run(parts[1])
                p_body.font.name = 'Times New Roman'
                p_body.font.size = Pt(12)
        else:
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
        return p

    def add_chapter_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(44, 82, 130) # Professional Steel Blue
        return p

    def add_page_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(44, 82, 130)
        return p

    def add_heading_styled(text, level):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(44, 82, 130)
        elif level == 2:
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(44, 82, 130)
        else:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_screenshot_placeholder(caption):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        set_cell_background(cell, "F7FAFC")
        set_cell_margins(cell, top=1000, bottom=1000, left=1000, right=1000)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run("\n\n[ INSERT MONITORING TELEMETRY SCREENSHOT HERE ]\n\n")
        run.font.name = 'Times New Roman'
        run.font.italic = True
        run.font.color.rgb = RGBColor(160, 174, 192)
        
        caption_p = doc.add_paragraph()
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_p.paragraph_format.space_before = Pt(4)
        caption_p.paragraph_format.space_after = Pt(12)
        run_cap = caption_p.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(10)
        run_cap.font.italic = True

    # --- COVER PAGE ---
    cover_p = doc.add_paragraph()
    cover_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_p.paragraph_format.space_before = Pt(48)
    cover_p.paragraph_format.space_after = Pt(12)
    c_run1 = cover_p.add_run("LOVELY PROFESSIONAL UNIVERSITY\n\n")
    c_run1.font.bold = True
    c_run1.font.size = Pt(16)
    c_run1.font.color.rgb = RGBColor(44, 82, 130)
    
    c_run2 = cover_p.add_run("SCHOOL OF COMPUTER SCIENCE AND ENGINEERING\n\n\n\n")
    c_run2.font.bold = True
    c_run2.font.size = Pt(13)
    
    c_run3 = cover_p.add_run(
        "FINAL INTERNSHIP REPORT\n\n"
        "CLOUD INFRASTRUCTURE MONITORING, SECURITY AUDITING\n"
        "AND AUTOMATED ALERTING IN MULTI-CLOUD ENVIRONMENTS\n\n\n"
    )
    c_run3.font.bold = True
    c_run3.font.size = Pt(16)
    c_run3.font.color.rgb = RGBColor(44, 82, 130)
    
    details_p = doc.add_paragraph()
    details_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    details_p.paragraph_format.line_spacing = 1.3
    det_run = details_p.add_run(
        "An Internship Report Completed at WiSys, Riyadh, Saudi Arabia\n\n"
        "Submitted in partial fulfillment of the requirements for the award of the degree of\n"
        "BACHELOR OF TECHNOLOGY\n"
        "IN\n"
        "COMPUTER SCIENCE AND ENGINEERING\n"
        "Specialization/Minor: Cloud Computing\n\n"
        "Course Code: CSE443 | Term: 26271\n\n\n"
        "SUBMITTED BY\n"
        "AQUIB JAWAID ANSARI\n"
        "Registration Number: 12508688\n"
        "Section: [SECTION]\n\n"
        "INDUSTRY MENTOR: Emadeldin Taha Mahrous\n"
        "COMPANY SUPERVISOR: Medhat\n"
        "FACULTY SUPERVISOR: [FACULTY SUPERVISOR]\n\n"
        "INTERNSHIP PERIOD: 14 June 2026 to 13 August 2026\n\n\n"
        "LOVELY PROFESSIONAL UNIVERSITY\n"
        "PHAGWARA, PUNJAB\n"
        "AUGUST 2026\n"
    )
    det_run.font.size = Pt(12)
    
    doc.add_page_break()
    
    # --- CERTIFICATE PAGE ---
    add_page_title("INTERNSHIP CERTIFICATE")
    p_cert = doc.add_paragraph()
    p_cert.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cert.paragraph_format.line_spacing = 1.5
    run = p_cert.add_run(
        "This page is reserved for the official internship certificate issued by WiSys. The certificate "
        "should confirm that Aquib Jawaid Ansari, Registration Number 12508688, successfully completed industry "
        "training at WiSys, Riyadh, Saudi Arabia, from 14 June 2026 to 13 August 2026.\n\n\n"
        "[INSERT SCANNED OFFICIAL INTERNSHIP CERTIFICATE IMAGE HERE]\n\n\n"
    )
    run.font.size = Pt(12)
    run.font.italic = True
    
    # Table of certificate details
    c_table = doc.add_table(rows=1, cols=2)
    c_table.style = 'Table Grid'
    hdr = c_table.rows[0].cells
    hdr[0].text = "Field"
    hdr[1].text = "Details"
    set_cell_background(hdr[0], "2C5282")
    set_cell_background(hdr[1], "2C5282")
    hdr[0].paragraphs[0].runs[0].font.bold = True
    hdr[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    hdr[1].paragraphs[0].runs[0].font.bold = True
    hdr[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    c_data = [
        ("Industry Mentor", "Emadeldin Taha Mahrous"),
        ("Company Supervisor", "Medhat"),
        ("Organization", "WiSys, Riyadh, Saudi Arabia")
    ]
    for k, v in c_data:
        row = c_table.add_row().cells
        row[0].text = k
        row[1].text = v
        set_cell_margins(row[0])
        set_cell_margins(row[1])
        
    doc.add_page_break()
    
    # --- DECLARATION PAGE ---
    add_page_title("STUDENT DECLARATION")
    add_standard_p(
        "I hereby declare that this internship report titled “Cloud Infrastructure Monitoring, Security Auditing and "
        "Automated Alerting in Multi-Cloud Environments” is an authentic record of the work performed by me during "
        "my internship at WiSys, Riyadh, Saudi Arabia, from 14 June 2026 to 13 August 2026."
    )
    add_standard_p(
        "The report describes the technical activities, learning outcomes, troubleshooting work and academic project "
        "developed from the knowledge gained during the internship. No confidential company credentials, passwords, "
        "access keys, private customer information or restricted production information have been included in this report."
    )
    add_standard_p(
        "This report is submitted in partial fulfillment of the requirements of CSE443 at Lovely Professional University.\n\n\n"
        "Date: 20 August 2026\n\n"
        "Student Signature: ______________________________\n\n"
        "Name: Aquib Jawaid Ansari\n"
        "Registration Number: 12508688"
    )
    
    doc.add_page_break()
    
    # --- ACKNOWLEDGEMENT PAGE ---
    add_page_title("ACKNOWLEDGEMENT")
    add_standard_p(
        "I sincerely thank WiSys, Riyadh, Saudi Arabia, for providing me with the opportunity to gain practical "
        "experience in cloud computing, infrastructure monitoring and IT operations."
    )
    add_standard_p(
        "I am especially thankful to my Industry Mentor, Emadeldin Taha Mahrous, and my Company Supervisor, Medhat, "
        "for their technical guidance, support and encouragement throughout the internship."
    )
    add_standard_p(
        "I would also like to thank the Cloud Operations and IT Infrastructure team for allowing me to work on practical "
        "tasks involving cloud monitoring, security auditing, automated alerting, troubleshooting and documentation."
    )
    add_standard_p(
        "I am also thankful to Lovely Professional University and the School of Computer Science and Engineering for "
        "providing the academic framework for this internship under course CSE443."
    )
    add_standard_p(
        "Finally, I would like to thank everyone who supported and guided me during the internship and helped me improve "
        "my technical and professional skills."
    )
    
    doc.add_page_break()
    
    # --- ABSTRACT PAGE ---
    add_page_title("ABSTRACT")
    add_standard_p(
        "Modern organizations depend heavily on cloud infrastructure for hosting applications, databases, networks "
        "and business services. As cloud environments become larger and more distributed, continuous monitoring, security "
        "auditing and automated alerting become essential for maintaining system availability, reliability and security."
    )
    add_standard_p(
        "This report presents the work completed during an internship at WiSys, Riyadh, Saudi Arabia, from 14 June 2026 "
        "to 13 August 2026. The work focused on Google Cloud Platform (GCP) and Huawei Cloud, with practical exposure to "
        "cloud infrastructure monitoring, security auditing, alert configuration, troubleshooting and automation. The "
        "internship involved monitoring approximately 40 virtual machines and related cloud resources, with attention to "
        "VM lifecycle events, health-check failures, high disk utilization, security-group modifications, VPC deletion, "
        "backup failures, audit-log changes and monitoring-policy changes."
    )
    add_standard_p(
        "In GCP, Cloud Logging, Cloud Monitoring and Cloud Audit Logs were used to understand and monitor infrastructure "
        "operations. Log-based custom metrics were applied to important VM lifecycle events, and security findings were "
        "reviewed and documented in a structured report. Another major contribution was the improvement of alert emails "
        "by adding available incident details such as instance name, IP address, event type, user identity and source IP, "
        "enabling administrators to perform an initial assessment without immediately opening the cloud console."
    )
    add_standard_p(
        "Huawei Cloud exposure included Cloud Trace Service (CTS), Cloud Eye, Simple Message Notification (SMN), Elastic "
        "Cloud Server (ECS) and Virtual Private Cloud (VPC). Troubleshooting work covered SMTP sender validation, alert "
        "deduplication, stale alert records, JSON configuration portability and event simulation. The internship also "
        "included Microsoft Entra ID configuration for the WiSys AI Hub application and completion of an AWS certification."
    )
    add_standard_p(
        "After the internship ended and access to the company cloud environment was no longer available, a standalone "
        "Multi-Cloud Monitoring Simulation was developed for academic evaluation. The simulation demonstrates resources, "
        "infrastructure events, alert rules, alert processing, alert deduplication and customized notifications in a "
        "controlled, non-production environment.\n\n"
        "Keywords: Cloud Monitoring, GCP, Huawei Cloud, Security Auditing, Automated Alerting, Cloud Audit Logs, "
        "Multi-Cloud, Microsoft Entra ID."
    )
    
    doc.add_page_break()
    
    # --- TABLE OF CONTENTS AS A REAL TABLE ---
    add_page_title("TABLE OF CONTENTS")
    toc_table = doc.add_table(rows=1, cols=3)
    toc_table.style = 'Table Grid'
    t_hdr = toc_table.rows[0].cells
    t_headers = ['No.', 'Chapter / Section', 'Page']
    for i, h in enumerate(t_headers):
        t_hdr[i].text = h
        set_cell_background(t_hdr[i], "2C5282")
        t_hdr[i].paragraphs[0].runs[0].font.bold = True
        t_hdr[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(t_hdr[i])
        
    toc_rows = [
        ("1", "Introduction", "1"),
        ("2", "Organization Profile", "2"),
        ("3", "Internship Role and Objectives", "3"),
        ("4", "Technologies and Cloud Environment", "4"),
        ("5", "Cloud Infrastructure Monitoring", "5"),
        ("6", "GCP Monitoring and Security Auditing", "6"),
        ("7", "Automated Alerting and Customized Notifications", "7"),
        ("8", "Huawei Cloud Monitoring", "8"),
        ("9", "Security Alert Configuration", "9"),
        ("10", "Work Performed During the Internship", "11"),
        ("11", "Major Troubleshooting Activities", "13"),
        ("12", "WiSys AI Hub and Microsoft Entra ID", "15"),
        ("13", "AWS Certification", "16"),
        ("14", "Multi-Cloud Monitoring Simulation", "17"),
        ("15", "Testing and Validation", "21"),
        ("16", "Learning Outcomes", "23"),
        ("17", "Challenges and Solutions", "24"),
        ("18", "Conclusion", "25"),
        ("19", "Future Scope", "26"),
        ("20", "References", "27"),
        ("21", "Appendices", "28")
    ]
    for num, name, pg in toc_rows:
        row_cells = toc_table.add_row().cells
        row_cells[0].text = num
        row_cells[1].text = name
        row_cells[2].text = pg
        for cell in row_cells:
            set_cell_margins(cell)
            
    doc.add_page_break()
    
    # --- LIST OF TABLES & FIGURES & ABBREVIATIONS ---
    list_t = doc.add_paragraph()
    lt_run = list_t.add_run("LIST OF TABLES\n")
    lt_run.font.bold = True
    list_t.add_run(
        "1. Table 1. Student and Internship Details\n"
        "2. Table 2. Technologies and Cloud Services Used\n"
        "3. Table 3. Core Monitoring and Security Alerts\n"
        "4. Table 4. Major Internship Contributions\n"
        "5. Table 5. Simulated Cloud Resources\n"
        "6. Table 6. Testing and Validation Results\n"
        "7. Table 7. Challenge and Solution Log\n"
    )
    
    list_f = doc.add_paragraph()
    lf_run = list_f.add_run("\nLIST OF FIGURES\n")
    lf_run.font.bold = True
    list_f.add_run(
        "8. Figure 1. Conceptual Multi-Cloud Monitoring Architecture\n"
        "9. Figure 2. AWS Certification Completed During Internship\n"
        "10. Figure 3. Multi-Cloud Monitoring Simulation Dashboard\n"
        "11. Figure 4. Simulated Infrastructure Alert\n"
        "12. Figure 5. Customized Alert Notification\n"
    )
    
    abbr_p = doc.add_paragraph()
    abbr_run = abbr_p.add_run("\nLIST OF ABBREVIATIONS\n")
    abbr_run.font.bold = True
    abbr_p.add_run(
        "GCP\t\tGoogle Cloud Platform\n"
        "AWS\t\tAmazon Web Services\n"
        "VM\t\tVirtual Machine\n"
        "ECS\t\tElastic Cloud Server\n"
        "CTS\t\tCloud Trace Service\n"
        "SMN\t\tSimple Message Notification\n"
        "VPC\t\tVirtual Private Cloud\n"
        "IAM\t\tIdentity and Access Management\n"
        "SMTP\t\tSimple Mail Transfer Protocol\n"
        "API\t\tApplication Programming Interface\n"
        "JSON\t\tJavaScript Object Notation\n"
        "IP\t\tInternet Protocol\n"
        "SLA\t\tService Level Agreement\n"
        "MTTR\t\tMean Time to Resolution\n"
        "LPU\t\tLovely Professional University\n"
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 1 ---
    add_chapter_title("1. INTRODUCTION")
    add_standard_p(
        "1.1 Background of the Internship\n"
        "I completed my internship at WiSys, Riyadh, Saudi Arabia, from 14 June 2026 to 13 August 2026. The main focus "
        "of the internship was cloud infrastructure monitoring, security auditing, automated alerting and IT infrastructure "
        "support. During the internship, I worked with cloud resources hosted on Google Cloud Platform and Huawei Cloud. "
        "My responsibilities included monitoring infrastructure, configuring alerts, analysing logs, reviewing security findings, "
        "troubleshooting technical issues and preparing technical documentation."
    )
    
    # Insert Table 1. Student and Internship Details
    add_heading_styled("Table 1. Student and Internship Details", 3)
    t1 = doc.add_table(rows=1, cols=2)
    t1.style = 'Table Grid'
    t1_hdr = t1.rows[0].cells
    t1_hdr[0].text = "Item"
    t1_hdr[1].text = "Details"
    set_cell_background(t1_hdr[0], "2C5282")
    set_cell_background(t1_hdr[1], "2C5282")
    t1_hdr[0].paragraphs[0].runs[0].font.bold = True
    t1_hdr[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    t1_hdr[1].paragraphs[0].runs[0].font.bold = True
    t1_hdr[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    t1_data = [
        ("Student Name", "Aquib Jawaid Ansari"),
        ("Registration Number", "12508688"),
        ("Programme", "B.Tech Computer Science and Engineering"),
        ("Specialization / Minor", "Cloud Computing"),
        ("Course Code / Term", "CSE443 / 26271"),
        ("Organization", "WiSys, Riyadh, Saudi Arabia"),
        ("Internship Period", "14 June 2026 to 13 August 2026"),
        ("Industry Mentor", "Emadeldin Taha Mahrous"),
        ("Company Supervisor", "Medhat"),
        ("Faculty Supervisor", "[FACULTY SUPERVISOR]")
    ]
    for k, v in t1_data:
        row = t1.add_row().cells
        row[0].text = k
        row[1].text = v
        set_cell_margins(row[0])
        set_cell_margins(row[1])
        
    add_standard_p("")
    add_standard_p(
        "1.2 Importance of Cloud Computing\n"
        "Cloud computing enables organizations to use computing resources such as virtual machines, storage, networks and "
        "applications without depending only on physical infrastructure. Cloud platforms make it possible to create, modify "
        "and manage infrastructure quickly; however, larger environments require disciplined monitoring because a single "
        "unexpected change can affect applications and users."
    )
    add_standard_p(
        "1.3 Importance of Cloud Monitoring\n"
        "Cloud monitoring helps administrators identify unexpected VM shutdowns, restarts and deletions; health-check "
        "failures; high disk utilization; network changes; backup failures; security changes and monitoring failures. "
        "Automated alerts allow the responsible team to receive important information quickly."
    )
    add_standard_p(
        "1.4 Importance of Security Auditing\n"
        "Cloud audit logs record administrative and system activities. Depending on the event, the available information "
        "may include user identity, source IP, resource name, operation, project, event time and operation details. "
        "This information supports both monitoring and security investigation."
    )
    add_standard_p(
        "1.5 Academic Relevance\n"
        "The internship was directly related to the B.Tech Computer Science and Engineering programme and Cloud Computing "
        "specialization. It enabled the practical application of cloud computing, computer networks, operating systems, "
        "cybersecurity, Python programming, database concepts, system administration, software engineering and cloud architecture."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 2 ---
    add_chapter_title("2. ORGANIZATION PROFILE")
    add_standard_p(
        "2.1 WiSys\n"
        "WiSys is a technology and IT services organization based in Riyadh, Saudi Arabia. The organization works in "
        "areas related to technology integration, cloud infrastructure, IT operations, cybersecurity and digital "
        "solutions. During the internship, the work was associated with cloud operations and IT infrastructure support."
    )
    add_standard_p(
        "2.2 Cloud Operations\n"
        "Cloud operations involve managing and monitoring infrastructure to maintain service availability and security. "
        "A monitoring system helps an organization detect incidents, identify affected resources, understand what happened, "
        "notify administrators, reduce response time and maintain availability. These activities formed an important "
        "part of the internship."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 3 ---
    add_chapter_title("3. INTERNSHIP ROLE AND OBJECTIVES")
    add_standard_p(
        "3.1 My Role\n"
        "I worked as a Cloud / IT Infrastructure Trainee. The role combined infrastructure observation, alert configuration, "
        "troubleshooting, automation and documentation. Key activities included:\n"
        "- Cloud infrastructure monitoring\n"
        "- Alert configuration\n"
        "- GCP log analysis\n"
        "- GCP security findings review\n"
        "- Security report preparation\n"
        "- Huawei Cloud monitoring\n"
        "- Audit-log analysis\n"
        "- SMTP troubleshooting\n"
        "- Python automation\n"
        "- Database-related troubleshooting\n"
        "- JSON configuration work\n"
        "- Microsoft Entra ID configuration\n"
        "- Technical documentation"
    )
    add_standard_p(
        "3.2 Main Objectives\n"
        "1. Understand enterprise cloud infrastructure.\n"
        "2. Learn practical cloud monitoring.\n"
        "3. Configure infrastructure alerts.\n"
        "4. Understand cloud audit logs.\n"
        "5. Identify and document security findings.\n"
        "6. Improve automated email notifications.\n"
        "7. Troubleshoot cloud-related problems.\n"
        "8. Automate repetitive technical tasks.\n"
        "9. Develop practical cloud engineering skills."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 4 ---
    add_chapter_title("4. TECHNOLOGIES AND CLOUD ENVIRONMENT")
    add_standard_p(
        "The table below details the technologies, platforms, and services used during the internship training:"
    )
    
    # Table 2: Technologies and Cloud Services Used
    add_heading_styled("Table 2. Technologies and Cloud Services Used", 3)
    t2 = doc.add_table(rows=1, cols=3)
    t2.style = 'Table Grid'
    t2_hdr = t2.rows[0].cells
    t2_hdr[0].text = "Technology / Service"
    t2_hdr[1].text = "Purpose"
    t2_hdr[2].text = "Exposure"
    for cell in t2_hdr:
        set_cell_background(cell, "2C5282")
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell)
        
    t2_data = [
        ("Google Cloud Platform", "Cloud infrastructure", "Hands-on"),
        ("Compute Engine", "Virtual machines", "Hands-on"),
        ("Cloud Logging", "Log analysis", "Hands-on"),
        ("Cloud Monitoring", "Monitoring and alerts", "Hands-on"),
        ("Cloud Audit Logs", "Activity auditing", "Hands-on"),
        ("Huawei Cloud", "Cloud infrastructure", "Hands-on"),
        ("Cloud Trace Service", "Audit events", "Hands-on"),
        ("Cloud Eye", "Monitoring and alarms", "Hands-on"),
        ("SMN", "Notifications", "Hands-on"),
        ("Microsoft Entra ID", "Application identity", "Hands-on"),
        ("Brevo SMTP", "Email delivery", "Hands-on"),
        ("Python", "Automation", "Hands-on"),
        ("Linux / Bash", "Server administration", "Practical"),
        ("JSON", "Configuration portability", "Hands-on"),
        ("AWS", "Cloud knowledge", "Certification")
    ]
    for tech, purp, exp in t2_data:
        row = t2.add_row().cells
        row[0].text = tech
        row[1].text = purp
        row[2].text = exp
        set_cell_margins(row[0])
        set_cell_margins(row[1])
        set_cell_margins(row[2])
        
    doc.add_page_break()
    
    # --- CHAPTER 5 ---
    add_chapter_title("5. CLOUD INFRASTRUCTURE MONITORING")
    add_standard_p(
        "5.1 Multi-Cloud Monitoring\n"
        "The internship involved monitoring activities across GCP and Huawei Cloud. Approximately 40 virtual machines "
        "and related resources were monitored across different cloud accounts and projects. The general monitoring process "
        "involved collecting information from cloud resources, analysing logs or metrics, applying monitoring rules, "
        "generating alerts and sending notifications to administrators."
    )
    add_standard_p(
        "5.2 VM Lifecycle Monitoring\n"
        "Important VM lifecycle events included VM stopped, VM restarted and VM deleted. These events are significant "
        "because unexpected changes can cause service interruption or indicate an unauthorized action."
    )
    add_standard_p(
        "5.3 Health Monitoring\n"
        "Health monitoring was used to identify whether important services or ports were responding correctly. A failed "
        "health check can indicate an application failure, network problem, port problem, VM problem or service failure."
    )
    add_standard_p(
        "5.4 Storage Monitoring\n"
        "Disk utilization monitoring was configured around a 90% threshold. High disk usage can result in application "
        "failures, database write problems, logging failures and service instability."
    )
    add_standard_p(
        "5.5 Network Monitoring\n"
        "Network-related changes were monitored, including security-group modifications and VPC deletion. These events "
        "may affect both security and availability."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 6 ---
    add_chapter_title("6. GCP MONITORING AND SECURITY AUDITING")
    add_standard_p(
        "6.1 GCP Monitoring\n"
        "The main GCP services used were Compute Engine, Cloud Logging, Cloud Monitoring and Cloud Audit Logs. One key "
        "task was monitoring VM lifecycle operations."
    )
    add_standard_p(
        "6.2 Log-Based Custom Metrics\n"
        "Some infrastructure operations are better identified through audit logs rather than normal resource metrics. "
        "Log filters were used to identify important VM lifecycle events and convert relevant log events into monitoring "
        "signals, allowing infrastructure events to act as inputs for monitoring and alert generation."
    )
    add_standard_p(
        "6.3 Event Information\n"
        "Depending on the event, audit information could include principal email, source IP, resource name, operation, "
        "project, event time and event details. This information was useful for generating customized alert notifications."
    )
    add_standard_p(
        "6.4 GCP Security Findings\n"
        "Security findings in GCP were reviewed and organized into a security findings report. The report recorded "
        "the finding type, affected resource, security importance, available evidence, recommended action and status."
    )
    add_standard_p(
        "6.5 Security Findings Reporting Process\n"
        "The process involved reviewing a security finding, identifying the affected resource, analysing the potential "
        "risk, documenting the issue and recommending an appropriate action. This work provided practical exposure "
        "to cloud security monitoring and security documentation."
    )
    add_standard_p(
        "6.6 Importance of the Security Review\n"
        "Security findings can reveal configuration problems or potential risks. Reviewing and documenting findings "
        "helps administrators understand what the problem is, which resource is affected, why it is important and what "
        "action should be taken."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 7 ---
    add_chapter_title("7. AUTOMATED ALERTING AND CUSTOMIZED NOTIFICATIONS")
    add_standard_p(
        "7.1 Basic Alerting Problem\n"
        "A practical issue with basic alert notifications is that the email may not contain enough information for an "
        "administrator to immediately understand an incident. Searching the cloud console for an affected resource "
        "can increase investigation time."
    )
    add_standard_p(
        "7.2 Custom Alert Messages\n"
        "Alert messages were improved by including useful incident information directly in the email. Depending on the "
        "event, the message could contain instance name, instance ID, IP address, event type, user or principal, "
        "source IP, project, resource, event time, severity and a short description."
    )
    
    add_standard_p("7.3 Sample Custom Alert Structure:")
    # Sample Custom Alert Table
    t_sa = doc.add_table(rows=1, cols=2)
    t_sa.style = 'Table Grid'
    sa_hdr = t_sa.rows[0].cells
    sa_hdr[0].text = "Field"
    sa_hdr[1].text = "Sample Value"
    set_cell_background(sa_hdr[0], "2C5282")
    set_cell_background(sa_hdr[1], "2C5282")
    sa_hdr[0].paragraphs[0].runs[0].font.bold = True
    sa_hdr[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    sa_hdr[1].paragraphs[0].runs[0].font.bold = True
    sa_hdr[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    sa_data = [
        ("ALERT TYPE", "VM INSTANCE STOPPED"),
        ("CLOUD PROVIDER", "GCP"),
        ("INSTANCE", "gcp-vm-prod-01"),
        ("IP ADDRESS", "192.168.x.x"),
        ("USER", "administrator"),
        ("SOURCE IP", "10.x.x.x"),
        ("SEVERITY", "CRITICAL"),
        ("EVENT TIME", "[EVENT TIME]"),
        ("DESCRIPTION", "The monitored virtual machine was stopped."),
        ("RECOMMENDED ACTION", "Verify whether the operation was authorized and check the application status.")
    ]
    for k, v in sa_data:
        row = t_sa.add_row().cells
        row[0].text = k
        row[1].text = v
        set_cell_margins(row[0])
        set_cell_margins(row[1])
        
    add_standard_p("")
    add_standard_p(
        "7.4 Advantage of Custom Alert Emails\n"
        "Customized messages allow an administrator to understand the basic incident directly from the email: what happened, "
        "which instance was affected, who performed the action, what source IP was involved and when the event occurred. "
        "This reduces the need to immediately open multiple cloud-console pages."
    )
    add_standard_p(
        "7.5 Notification Process\n"
        "The notification process involved detecting a cloud event, processing the corresponding log or metric, "
        "evaluating the alert rule, generating a customized message, sending it through the SMTP relay and delivering "
        "the notification to the administrator."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 8 ---
    add_chapter_title("8. HUAWEI CLOUD MONITORING")
    add_standard_p(
        "8.1 Huawei Cloud\n"
        "Huawei Cloud monitoring work involved Cloud Trace Service (CTS), Cloud Eye, Simple Message Notification (SMN), "
        "Elastic Cloud Server (ECS) and Virtual Private Cloud (VPC)."
    )
    add_standard_p(
        "8.2 Cloud Trace Service\n"
        "Cloud Trace Service records important operations performed on Huawei Cloud resources. It was used to understand "
        "administrative events and resource changes, with event details helping identify the resource, operation, user, "
        "time and result."
    )
    add_standard_p(
        "8.3 Cloud Eye\n"
        "Cloud Eye was used for monitoring and alarm management. Work included alarm configuration and understanding the "
        "relationship between events, monitoring rules and notifications."
    )
    add_standard_p(
        "8.4 Simple Message Notification\n"
        "SMN was used to route notifications through topics. The flow consisted of event detection, rule evaluation, "
        "Cloud Eye alarm generation and notification forwarding through an SMN topic."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 9 ---
    add_chapter_title("9. SECURITY ALERT CONFIGURATION")
    add_standard_p(
        "The table below details the 11 monitored alerts configured during my internship:"
    )
    
    # Table 3: Core Monitoring and Security Alerts
    add_heading_styled("Table 3. Core Monitoring and Security Alerts", 3)
    t3 = doc.add_table(rows=1, cols=4)
    t3.style = 'Table Grid'
    t3_hdr = t3.rows[0].cells
    t3_headers = ['Category', 'Alert', 'Severity', 'Main Risk']
    for i, h in enumerate(t3_headers):
        t3_hdr[i].text = h
        set_cell_background(t3_hdr[i], "2C5282")
        t3_hdr[i].paragraphs[0].runs[0].font.bold = True
        t3_hdr[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(t3_hdr[i])
        
    t3_data = [
        ("Compute", "VM Instance Deleted", "Critical", "Loss of resource or system state"),
        ("Compute", "VM Instance Stopped", "Critical", "Application downtime"),
        ("Compute", "VM Instance Restarted", "Critical", "Unexpected service interruption"),
        ("Compute", "VM Health Check Failed", "Critical", "Service unavailable"),
        ("Storage", "Disk Utilization \u2265 90%", "Critical", "Storage exhaustion"),
        ("Network", "Security Group Modified", "Critical", "Possible security exposure"),
        ("Network", "VPC Deleted", "Critical", "Network disruption"),
        ("Backup", "Backup Failed", "Critical", "Data-loss risk"),
        ("Security", "Audit Log Disabled", "Critical", "Loss of audit visibility"),
        ("Monitoring", "Alarm Policy Deleted", "Critical", "Monitoring gap"),
        ("Monitoring", "Alarm Policy Disabled", "Warning", "Resource may become unmonitored")
    ]
    for cat, name, sev, risk in t3_data:
        row = t3.add_row().cells
        row[0].text = cat
        row[1].text = name
        row[2].text = sev
        row[3].text = risk
        set_cell_margins(row[0])
        set_cell_margins(row[1])
        set_cell_margins(row[2])
        set_cell_margins(row[3])
        
    add_standard_p("")
    add_standard_p(
        "9.1 VM Instance Deleted\n"
        "A VM deletion event can result in loss of the VM and its system state; it is therefore an important event to monitor."
    )
    add_standard_p(
        "9.2 VM Instance Stopped\n"
        "An unexpected VM stop can cause application downtime. The alert allows administrators to identify the event "
        "quickly and investigate the reason."
    )
    add_standard_p(
        "9.3 VM Instance Restarted\n"
        "An unexpected restart can indicate system failure, a software issue, an administrative action or an "
        "infrastructure problem."
    )
    add_standard_p(
        "9.4 VM Health Check Failed\n"
        "A failed health check indicates that a service or system is not responding correctly."
    )
    add_standard_p(
        "9.5 Disk Utilization \u2265 90%\n"
        "High disk usage can cause application and database problems; the 90% threshold was treated as an important alert condition."
    )
    add_standard_p(
        "9.6 Security Group Modified\n"
        "Security-group modifications can change network access. An unexpected change can create a security risk."
    )
    add_standard_p(
        "9.7 VPC Deleted\n"
        "VPC deletion can affect connected resources and network connectivity, and should be detected immediately."
    )
    add_standard_p(
        "9.8 Backup Failed\n"
        "Backup failure increases the risk of data loss. Monitoring backup failures allows administrators to take "
        "corrective action."
    )
    add_standard_p(
        "9.9 Audit Log Disabled\n"
        "Audit logs are important for security investigations. Disabling logging reduces visibility into administrative actions."
    )
    add_standard_p(
        "9.10 Alarm Policy Deleted or Disabled\n"
        "Deleting or disabling monitoring policies can create monitoring gaps. These changes should therefore be "
        "detected and reviewed."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 10 ---
    add_chapter_title("10. WORK PERFORMED DURING THE INTERNSHIP")
    add_standard_p(
        "The table below summarizes my specific engineering and operational tasks during the training:"
    )
    
    # Table 4: Major Internship Contributions
    add_heading_styled("Table 4. Major Internship Contributions", 3)
    t4 = doc.add_table(rows=1, cols=2)
    t4.style = 'Table Grid'
    t4_hdr = t4.rows[0].cells
    t4_hdr[0].text = "Area"
    t4_hdr[1].text = "Contribution"
    set_cell_background(t4_hdr[0], "2C5282")
    set_cell_background(t4_hdr[1], "2C5282")
    t4_hdr[0].paragraphs[0].runs[0].font.bold = True
    t4_hdr[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    t4_hdr[1].paragraphs[0].runs[0].font.bold = True
    t4_hdr[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    t4_data = [
        ("Multi-Cloud Monitoring", "Configured and verified monitoring for approximately 40 virtual machines and related resources."),
        ("GCP Custom Metrics", "Configured log filters to track VM lifecycle operations."),
        ("GCP Security", "Reviewed security findings and prepared a security findings report."),
        ("GCP Alerts", "Configured and tested monitoring alerts."),
        ("Custom Email Alerts", "Added instance, IP, user, source and event information to notification messages."),
        ("Huawei Cloud", "Worked with CTS, Cloud Eye and SMN."),
        ("SMTP", "Troubleshot email delivery and sender validation problems."),
        ("Alert Database", "Identified and resolved a backlog of 55 stale OPEN alerts."),
        ("JSON Configuration", "Worked on reusable alarm configurations and project-specific parameters."),
        ("Event Simulator", "Improved the simulator to use valid resource information."),
        ("WiSys AI Hub", "Configured Microsoft Entra ID application registration and related settings."),
        ("Documentation", "Prepared technical documentation and troubleshooting information."),
        ("AWS", "Completed an AWS certification during the internship.")
    ]
    for area, contr in t4_data:
        row = t4.add_row().cells
        row[0].text = area
        row[1].text = contr
        set_cell_margins(row[0])
        set_cell_margins(row[1])
        
    add_standard_p("")
    add_standard_p(
        "10.1 Multi-Cloud Monitoring Setup\n"
        "Monitoring activities for approximately 40 virtual machines and related cloud resources were configured and "
        "verified to detect important infrastructure events and notify administrators."
    )
    add_standard_p(
        "10.2 GCP Custom Metrics\n"
        "Log filters on GCP Audit Logs were configured to identify VM lifecycle operations and convert relevant log "
        "events into monitoring signals."
    )
    add_standard_p(
        "10.3 GCP Security Findings\n"
        "GCP security findings were reviewed and documented in a security findings report, providing practical "
        "experience in identifying, analysing and documenting cloud security issues."
    )
    add_standard_p(
        "10.4 Custom Alert Emails\n"
        "Alert messages were customized to include useful incident information, such as instance details, IP address, "
        "user, source IP and event information."
    )
    add_standard_p(
        "10.5 Huawei Cloud Monitoring\n"
        "CTS, Cloud Eye and SMN were used to understand event monitoring and notification across the Huawei Cloud "
        "monitoring pipeline."
    )
    add_standard_p(
        "10.6 Technical Documentation\n"
        "Documentation was prepared for configurations, troubleshooting procedures and monitoring workflows, helping "
        "technical teams understand system configuration and common problem resolution."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 11 ---
    add_chapter_title("11. MAJOR TROUBLESHOOTING ACTIVITIES")
    add_standard_p(
        "11.1 SMTP Relay Sender Validation Problem\n"
        "During testing, SMTP alert emails were not delivered correctly because the SMTP service rejected the configured "
        "sender identity. Configuration and logs were investigated, the sender validation problem was identified and the "
        "email configuration was corrected. Email delivery was restored after correction."
    )
    add_standard_p(
        "11.2 Alert Deduplication Queue Problem\n"
        "The alerting system used deduplication logic to prevent repeated notifications for the same active alert. "
        "During testing, new alert emails were not generated. Database checks identified 55 stale OPEN alerts, which "
        "prevented new events from creating notifications. A Python cleanup script was created and executed to change "
        "stale alerts to RESOLVED. The stale backlog was removed, new alerts could be generated, notification testing "
        "worked correctly and the dashboard compliance score returned to 100/100."
    )
    add_standard_p(
        "11.3 JSON Configuration Import Problem\n"
        "Exported JSON alarm configurations failed to import in some cases because they included project-specific "
        "identifiers. These values were parameterized so configurations could be reused with different project values, "
        "improving portability."
    )
    add_standard_p(
        "11.4 Event Simulator Problem\n"
        "The event simulator initially used dummy resource IDs, meaning simulated events did not always update the "
        "correct dashboard resources. It was improved to use valid resource information from the database, making "
        "testing more realistic and associating simulated events with represented resources."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 12 ---
    add_chapter_title("12. WISYS AI HUB AND MICROSOFT ENTRA ID")
    add_standard_p(
        "12.1 WiSys AI Hub\n"
        "Identity configuration work was completed for the WiSys AI Hub application. The application required "
        "integration with Microsoft Entra ID for authentication."
    )
    add_standard_p(
        "12.2 Application Registration\n"
        "Specific configurations included:\n"
        "- Application registration in the Entra portal\n"
        "- Client credential (secret) generation\n"
        "- Redirect URI configuration\n"
        "- API permission configuration\n"
        "- Requesting User.Read permission\n"
        "- Acquiring Administrator consent"
    )
    add_standard_p(
        "12.3 Importance of the Configuration\n"
        "Microsoft Entra ID provides a centralized identity system for applications. The configuration enabled the "
        "WiSys AI Hub application to use standard authentication and authorization mechanisms, providing practical "
        "experience with application identity and access management."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 13 ---
    add_chapter_title("13. AWS CERTIFICATION")
    add_standard_p(
        "13.1 Certification Completed During the Internship\n"
        "An AWS certification was completed during the internship. It expanded cloud knowledge beyond the GCP and "
        "Huawei Cloud environments used in day-to-day internship work and strengthened understanding of cloud computing, "
        "AWS services, cloud infrastructure, cloud security, cloud architecture and cloud operations."
    )
    add_screenshot_placeholder("AWS Certification Completed During Internship")
    
    doc.add_page_break()
    
    # --- CHAPTER 14 ---
    add_chapter_title("14. MULTI-CLOUD MONITORING SIMULATION")
    add_standard_p(
        "14.1 Project Background\n"
        "After the internship ended, access to the company cloud environment was no longer available. Therefore, "
        "a standalone Multi-Cloud Monitoring Simulation was developed for academic evaluation and demonstration. "
        "The simulation is based on monitoring requirements and concepts learned during the internship, and it does "
        "not use WiSys credentials or company production infrastructure."
    )
    add_standard_p(
        "14.2 Project Objective\n"
        "The objective is to demonstrate how resources from multiple cloud providers can be represented and monitored "
        "using a common monitoring system. The simulation covers GCP and Huawei Cloud resources, VM lifecycle events, "
        "security events, storage alerts, network alerts, backup alerts, monitoring alerts and automated notifications."
    )
    add_standard_p(
        "14.3 Problem Statement\n"
        "In a multi-cloud environment, resources are distributed across providers. Monitoring each provider "
        "separately can make infrastructure management more difficult. The simulation provides a common monitoring "
        "environment in which resources and events from different providers can be represented together."
    )
    add_standard_p(
        "14.4 Proposed Architecture\n"
        "The proposed architecture comprises cloud resources, a Resource Manager, Event Generator, Monitoring Engine, "
        "Alert Rules, Alert Database, Notification Generator, Dashboard and Administrator. The conceptual sequence of "
        "monitoring and alerting components is represented in Figure 1."
    )
    
    # Insert conceptual architecture diagram
    add_screenshot_placeholder("Conceptual Multi-Cloud Monitoring Architecture (Flowchart)")
    
    add_standard_p(
        "14.5 Resource Manager\n"
        "The Resource Manager stores simulated resource information, including provider, resource ID, resource name, "
        "resource type, IP address, status and environment."
    )
    add_standard_p(
        "14.6 Event Generator\n"
        "The Event Generator creates simulated infrastructure events such as VM stopped, VM restarted, VM deleted, "
        "health check failed, disk utilization high, security group modified, VPC deleted, backup failed, audit "
        "logging disabled and alarm policy disabled."
    )
    add_standard_p(
        "14.7 Monitoring Engine\n"
        "The Monitoring Engine checks every generated event against configured alert rules. For example, when a VM stop "
        "event is generated, the engine identifies the VM Instance Stopped rule, assigns the appropriate severity "
        "and creates an alert."
    )
    add_standard_p(
        "14.8 Alert Database\n"
        "The alert database stores generated alerts and their status. Important states include OPEN and RESOLVED. "
        "It also supports alert deduplication so repeated events do not create unnecessary notifications."
    )
    add_standard_p(
        "14.9 Notification Generator\n"
        "The Notification Generator creates detailed messages containing available event information such as cloud "
        "provider, instance name, instance ID, IP address, event, user, source IP, time, severity and description."
    )
    
    # Table 5: Simulated Cloud Resources
    add_heading_styled("Table 5. Simulated Cloud Resources", 3)
    t5 = doc.add_table(rows=1, cols=4)
    t5.style = 'Table Grid'
    t5_hdr = t5.rows[0].cells
    t5_hdr[0].text = "Provider"
    t5_hdr[1].text = "Resource Name"
    t5_hdr[2].text = "Type"
    t5_hdr[3].text = "Status"
    for cell in t5_hdr:
        set_cell_background(cell, "2C5282")
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell)
        
    t5_data = [
        ("GCP", "gcp-vm-prod-01", "Compute VM", "RUNNING"),
        ("GCP", "gcp-vm-db-01", "Compute VM", "RUNNING"),
        ("Huawei Cloud", "huawei-ecs-prod-01", "ECS", "RUNNING"),
        ("Huawei Cloud", "huawei-ecs-db-01", "ECS", "RUNNING"),
        ("GCP", "gcp-vpc-prod", "VPC", "ACTIVE"),
        ("Huawei Cloud", "huawei-vpc-prod", "VPC", "ACTIVE")
    ]
    for prov, name, rtype, status in t5_data:
        row = t5.add_row().cells
        row[0].text = prov
        row[1].text = name
        row[2].text = rtype
        row[3].text = status
        set_cell_margins(row[0])
        set_cell_margins(row[1])
        set_cell_margins(row[2])
        set_cell_margins(row[3])
        
    add_standard_p("Note: These are simulated resources created for academic demonstration and are not company production resources.")
    
    add_standard_p(
        "14.10 Alert Deduplication\n"
        "If the same resource repeatedly generates the same alert while a previous alert remains OPEN, the system "
        "can avoid creating unnecessary duplicate alerts. When the issue is resolved, the alert is changed to RESOLVED, "
        "allowing a new event to create a new alert. This concept reflects the alert troubleshooting work completed "
        "during the internship."
    )
    add_standard_p(
        "14.11 Security Monitoring\n"
        "The simulation includes security-related alerts such as security group modification, audit logging disabled, "
        "VPC deletion and unauthorized-looking VM lifecycle events. Administrators can review the events through the "
        "dashboard and notification system."
    )
    add_standard_p(
        "14.12 Multi-Cloud Dashboard\n"
        "The dashboard provides a single view of simulated resources and alerts. It can display total resources, "
        "running resources, stopped resources, open alerts, critical alerts, resolved alerts, cloud provider and "
        "resource status."
    )
    add_screenshot_placeholder("Multi-Cloud Monitoring Simulation Dashboard")
    
    add_standard_p(
        "14.13 Simulation Benefits\n"
        "- Monitor different cloud providers through one conceptual interface.\n"
        "- Standardize alert rules and severities.\n"
        "- Display resources in one interface.\n"
        "- Generate detailed alerts and customized notifications.\n"
        "- Track alert status and reduce duplicate notifications.\n"
        "- Support security monitoring and a basis for future automation."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 15 ---
    add_chapter_title("15. TESTING AND VALIDATION")
    add_standard_p(
        "15.1 Testing Approach\n"
        "The monitoring simulation was tested by generating different infrastructure events and checking whether the "
        "correct alert was created. The tests verified event detection, alert type, severity, resource, status, "
        "notification information and deduplication behaviour."
    )
    
    # Table 6. Testing and Validation Results
    add_heading_styled("Table 6. Testing and Validation Results", 3)
    t6 = doc.add_table(rows=1, cols=3)
    t6.style = 'Table Grid'
    t6_hdr = t6.rows[0].cells
    t6_hdr[0].text = "Test Event"
    t6_hdr[1].text = "Expected Result"
    t6_hdr[2].text = "Status"
    for cell in t6_hdr:
        set_cell_background(cell, "2C5282")
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell)
        
    t6_data = [
        ("VM Stop", "VM Stopped alert", "Passed"),
        ("VM Restart", "VM Restarted alert", "Passed"),
        ("VM Delete", "VM Deleted alert", "Passed"),
        ("Health Check Failure", "Health alert", "Passed"),
        ("High Disk Usage", "Storage alert", "Passed"),
        ("Security Group Change", "Security alert", "Passed"),
        ("VPC Deletion", "Network alert", "Passed"),
        ("Backup Failure", "Backup alert", "Passed"),
        ("Audit Log Disabled", "Security alert", "Passed"),
        ("Alarm Policy Deleted", "Monitoring alert", "Passed"),
        ("Alarm Policy Disabled", "Monitoring warning", "Passed")
    ]
    for event, expected, status in t6_data:
        row = t6.add_row().cells
        row[0].text = event
        row[1].text = expected
        row[2].text = status
        set_cell_margins(row[0])
        set_cell_margins(row[1])
        set_cell_margins(row[2])
        
    add_standard_p("")
    add_standard_p(
        "15.2 Custom Email Test\n"
        "The custom notification was tested to verify inclusion of important fields such as event, provider, resource, "
        "IP, user, source IP, severity and time."
    )
    add_standard_p(
        "15.3 Deduplication Test\n"
        "Repeated events were generated for the same resource. The system checked whether an OPEN alert already existed. "
        "If an active alert existed, unnecessary duplicate alerts were avoided. After the existing alert was resolved, "
        "a new event could create a new alert."
    )
    add_screenshot_placeholder("Simulated Infrastructure Alert")
    add_screenshot_placeholder("Customized Alert Notification (Email)")
    
    doc.add_page_break()
    
    # --- CHAPTER 16 ---
    add_chapter_title("16. LEARNING OUTCOMES")
    add_standard_p(
        "16.1 Cloud Observability & Competencies\n"
        "- Cloud Computing: Practical understanding of cloud infrastructure, managing virtual machines, networks, and services.\n"
        "- Cloud Monitoring: Setting up monitoring rules, metrics, and alerting workflows.\n"
        "- Security Auditing: Utilizing cloud audit logs to track resource modifications.\n"
        "- Alert Automation: Configuring notification topics, subscriptions, and SMTP relays.\n"
        "- Troubleshooting: Solving sender validation rejections, clearing DB alert backlogs, and parameters portability.\n"
        "- Identity Management: Managing application registrations, client secrets, and permissions in Microsoft Entra ID."
    )
    add_standard_p(
        "16.2 Professional Skills\n"
        "The internship improved technical communication, ticket workflows, problem-solving, SLA compliance, "
        "and collaboration within a professional IT operations team."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 17 ---
    add_chapter_title("17. CHALLENGES AND SOLUTIONS")
    add_standard_p(
        "The table below details the core technical challenges and applied solutions during the internship:"
    )
    
    # Table 7: Challenge and Solution Log
    add_heading_styled("Table 7. Challenge and Solution Log", 3)
    t7 = doc.add_table(rows=1, cols=3)
    t7.style = 'Table Grid'
    t7_hdr = t7.rows[0].cells
    t7_hdr[0].text = "Challenge"
    t7_hdr[1].text = "Solution"
    t7_hdr[2].text = "Result"
    for cell in t7_hdr:
        set_cell_background(cell, "2C5282")
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell)
        
    t7_data = [
        ("SMTP sender rejection", "Corrected sender configuration to verified email address", "Email delivery restored"),
        ("55 stale OPEN alerts blocking queue", "Ran Python database cleanup script to transition states to RESOLVED", "New alerts processed correctly"),
        ("JSON import compatibility errors", "Parameterized environment-specific UUIDs and ARNs", "Rules standardized for multi-project reuse"),
        ("Dummy simulator resources", "Refactored generator loop to target real database resource entities", "Simulated events updated actual resources")
    ]
    for ch, sol, res in t7_data:
        row = t7.add_row().cells
        row[0].text = ch
        row[1].text = sol
        row[2].text = res
        set_cell_margins(row[0])
        set_cell_margins(row[1])
        set_cell_margins(row[2])
        
    doc.add_page_break()
    
    # --- CHAPTER 18 ---
    add_chapter_title("18. CONCLUSION")
    add_standard_p(
        "The internship at WiSys, Riyadh, provided valuable practical experience in cloud infrastructure, monitoring, "
        "security auditing and IT operations. Work with GCP and Huawei Cloud provided hands-on exposure to cloud monitoring, "
        "audit logs, custom metrics, security findings, alert configuration and automated notifications."
    )
    add_standard_p(
        "Improving alert emails with instance name, IP address, event type, user and source IP made it possible for "
        "administrators to perform an initial analysis directly from an email. GCP security findings work provided "
        "experience in cloud security assessment and reporting. Troubleshooting tasks involving SMTP delivery, alert "
        "deduplication, database records, JSON configurations and event simulation improved debugging and automation skills."
    )
    add_standard_p(
        "Microsoft Entra ID configuration for WiSys AI Hub provided experience with application identity, while AWS "
        "certification further expanded cloud knowledge. Finally, the Multi-Cloud Monitoring Simulation converted "
        "practical concepts learned at WiSys into an academic project that can be demonstrated without access to "
        "company cloud infrastructure."
    )
    add_standard_p(
        "Overall, the internship helped bridge the gap between academic learning and real-world cloud infrastructure "
        "operations and provided a strong foundation for further work in Cloud Engineering, DevOps, Infrastructure "
        "Engineering and Cloud Security."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 19 ---
    add_chapter_title("19. FUTURE SCOPE")
    add_standard_p(
        "1. Centralized multi-cloud log aggregation to view alerts from all cloud providers in a single console.\n"
        "2. Automating monitoring deployments using Infrastructure as Code (IaC) tools like Terraform.\n"
        "3. Implementing machine learning models to detect anomalies and unusual patterns in resource performance.\n"
        "4. Creating automated remediation playbooks to resolve incidents (such as restarting crashed ports) automatically.\n"
        "5. Integrating monitoring alerts with IT Service Management (ITSM) ticketing systems."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 20 ---
    add_chapter_title("20. REFERENCES")
    add_standard_p(
        "1. Google Cloud Documentation – Google Cloud Platform, Cloud Logging and Cloud Monitoring.\n"
        "2. Google Cloud Documentation – Cloud Audit Logs and Compute Engine.\n"
        "3. Huawei Cloud Documentation – Cloud Trace Service (CTS).\n"
        "4. Huawei Cloud Documentation – Cloud Eye Monitoring Service.\n"
        "5. Huawei Cloud Documentation – Simple Message Notification (SMN).\n"
        "6. Microsoft Documentation – Microsoft Entra ID Application Registration.\n"
        "7. Microsoft Documentation – Microsoft Entra ID API Permissions.\n"
        "8. Amazon Web Services Documentation – AWS Cloud Services and Architecture.\n"
        "9. Lovely Professional University – B.Tech Computer Science and Engineering Curriculum.\n"
        "10. WiSys Internship Training and Technical Work.\n\n"
        "Note: Company-internal material is referenced only at a high level. No confidential credentials, customer "
        "data or restricted production information is included in this report."
    )
    
    doc.add_page_break()
    
    # --- CHAPTER 21 ---
    add_chapter_title("21. APPENDICES")
    add_standard_p(
        "APPENDIX A – CONCEPTUAL ARCHITECTURE\n"
        "[GCP & Huawei Cloud Resources] ➔ [Resource Manager] ➔ [Event Generator] ➔ [Monitoring Engine & Rules] "
        "➔ [Alert DB & Deduplication] ➔ [Notification Engine & Dashboard] ➔ [Administrator]"
    )
    
    add_heading_styled("APPENDIX B – CUSTOM CONFIGURATION FORMATS", 2)
    p_code = doc.add_paragraph()
    p_code.paragraph_format.line_spacing = 1.2
    p_code.paragraph_format.space_before = Pt(4)
    p_code.paragraph_format.space_after = Pt(12)
    c_run = p_code.add_run(
        "Example 1: Parameterized JSON Config Template for SMTP Service (smtp_config.json)\n"
        "{\n"
        "  \"smtp_server\": \"smtp-relay.brevo.com\",\n"
        "  \"smtp_port\": 587,\n"
        "  \"smtp_username\": \"b5f116001@smtp-brevo.com\",\n"
        "  \"smtp_password\": \"[REDACTED_PASSWORD_KEY]\",\n"
        "  \"alert_recipient\": \"jawaidaquib893@gmail.com\"\n"
        "}\n\n"
        "Example 2: Sanitized GCP Custom Log Metric Filter\n"
        "resource.type=\"gce_instance\"\n"
        "protoPayload.methodName=\"v1.compute.instances.stop\"\n"
        "protoPayload.authenticationInfo.principalEmail=\"*@wisys.sa\"\n"
    )
    c_run.font.name = "Consolas"
    c_run.font.size = Pt(11)
    
    # Set page numbers
    add_header_footer(doc)
    
    # Save Document
    target_path = r"C:\Users\dell\Desktop\LPU\BTech 5th Sem\Summer Internship\Project\LPU_BTech_Final_Internship_Report.docx"
    doc.save(target_path)
    print(f"SUCCESS: Report saved as {target_path}")

if __name__ == "__main__":
    main()
